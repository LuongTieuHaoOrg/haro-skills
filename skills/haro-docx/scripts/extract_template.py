#!/usr/bin/env python3
"""Create a haro-docx template entry from a user-supplied .docx file.

Copies <source.docx> into the registry (local or global) and generates
template.yaml whose STYLE params are extracted from the .docx (style-level
only) merged over templates/config.yaml defaults, plus content.txt — a
machine dump of the .docx text/tables/headers for the AGENT to read and
understand the template's purpose. The script never guesses purpose or
content values; the agent proposes params and the user approves them.

Usage:
    python extract_template.py --source <file.docx> --id <template-id>
        --location local|global [--project-root .]
        [--name "Ten mau"] [--description "Mo ta"]

Language convention: code/comments in English; user-visible strings in
Vietnamese with full diacritics.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from template_store import (
    TEMPLATE_DOCX_NAME,
    TEMPLATE_YAML_NAME,
    copy_source_into,
    exists_in,
    global_root,
    local_root,
    normalize_id,
    now_iso,
    template_paths,
)

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULTS_YAML = SCRIPT_DIR / ".." / "templates" / "config.yaml"

CONTENT_TYPE_EXT = {
    "image/png": "png",
    "image/jpeg": "jpg",
    "image/jpg": "jpg",
    "image/gif": "gif",
    "image/bmp": "bmp",
}


def _num(value) -> float | int | None:
    """Normalize a Pt/cm number: int when whole, else rounded."""
    if value is None:
        return None
    try:
        num = round(float(value), 2)
    except (TypeError, ValueError):
        return None
    return int(num) if float(num).is_integer() else num


def _pt(length) -> float | int | None:
    if length is None:
        return None
    try:
        return _num(length.pt)
    except Exception:
        return None


def _cm(length) -> float | int | None:
    if length is None:
        return None
    try:
        return _num(length.cm)
    except Exception:
        return None


def _line_spacing(pf) -> float | None:
    """Return scalar line-spacing factor, or None for exact/at-least rules."""
    try:
        ls = pf.line_spacing
    except Exception:
        return None
    if ls is None:
        return None
    if isinstance(ls, float):
        return round(ls, 2)
    return None  # Length (exact rule) — not representable as a factor


def _style_font(doc: Document, name: str) -> dict:
    try:
        st = doc.styles[name]
    except KeyError:
        return {}
    out: dict = {}
    try:
        if st.font.name:
            out["name"] = st.font.name
    except Exception:
        pass
    size = _pt(getattr(st.font, "size", None))
    if size:
        out["size"] = size
    return out


def _style_paragraph(style) -> dict:
    """Extract spacing facts from a style's paragraph_format."""
    out: dict = {}
    try:
        pf = style.paragraph_format
    except Exception:
        return out
    ls = _line_spacing(pf)
    if ls is not None:
        out["line_spacing"] = ls
    sb = _pt(getattr(pf, "space_before", None))
    if sb is not None:
        out["space_before"] = sb
    sa = _pt(getattr(pf, "space_after", None))
    if sa is not None:
        out["space_after"] = sa
    indent = _cm(getattr(pf, "left_indent", None))
    if indent:
        out["indent_cm"] = indent
    return out


def _has_images_in_element(element) -> bool:
    for el in element.iter():
        tag = str(getattr(el, "tag", ""))
        if tag.endswith("}drawing") or tag.endswith("}pict") or tag.endswith("}blip"):
            return True
    return False


def _para_text_br(p) -> str:
    """Paragraph text with explicit line breaks kept as ' / ' separators."""
    parts = []
    try:
        for el in p._p.iter():
            tag = el.tag
            if not isinstance(tag, str):
                continue
            if tag.endswith("}t"):
                parts.append(el.text or "")
            elif tag.endswith("}br"):
                parts.append(" / ")
    except Exception:
        return (p.text or "").strip()
    text = "".join(parts).strip()
    return text if text else (p.text or "").strip()


def _part_text(part) -> str:
    texts = []
    for p in getattr(part, "paragraphs", []):
        text = _para_text_br(p)
        if text:
            texts.append(text)
    for table in getattr(part, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    text = _para_text_br(p)
                    if text:
                        texts.append(text)
    return " | ".join(texts[:6])


def _header_image_blobs(doc: Document) -> list[tuple[bytes, str]]:
    """Collect (blob, content_type) images from all section headers."""
    blobs = []
    seen = set()
    for sec in doc.sections:
        try:
            part = sec.header.part
        except Exception:
            continue
        for el in part._element.iter():
            if not str(getattr(el, "tag", "")).endswith("}blip"):
                continue
            rid = el.get(qn("r:embed"))
            if not rid or rid in seen:
                continue
            seen.add(rid)
            try:
                img = part.related_parts[rid]
                blobs.append((img.blob, img.content_type))
            except (KeyError, AttributeError):
                continue
    return blobs


def save_logo_blob(blobs: list[tuple[bytes, str]], dest_dir: Path) -> str:
    """Save the first header image under <template-dir>/assets/. Returns rel path or ''."""
    if not blobs:
        return ""
    blob, ctype = blobs[0]
    ext = CONTENT_TYPE_EXT.get((ctype or "").lower(), "png")
    assets = dest_dir / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    rel = f"assets/logo.{ext}"
    with open(dest_dir / rel, "wb") as f:
        f.write(blob)
    return rel


def dump_content(doc: Document, source_name: str, logo_rel: str = "",
                 max_paras: int = 60, max_rows: int = 20) -> str:
    """Dump readable .docx content for the AGENT to understand purpose.

    Plain facts only (texts, tables, headers/footers, image list) — no
    guessing, no confidence scores. Purpose inference happens in the
    agent workflow (commands/create.md), not here.
    """
    lines = [f"# Nội dung file mẫu: {source_name}",
             "# (dump máy để agent đọc hiểu mục đích — không phải gợi ý giá trị)",
             ""]
    paras = [p for p in doc.paragraphs if (p.text or "").strip()]
    lines.append(f"## Đoạn văn ({len(paras)} đoạn, hiện {min(len(paras), max_paras)} đầu)")
    for i, p in enumerate(paras[:max_paras]):
        tags = []
        try:
            if p.style and p.style.name not in ("Normal",):
                tags.append(p.style.name)
        except Exception:
            pass
        try:
            sizes = sorted({_pt(r.font.size) for r in p.runs
                            if getattr(r.font, "size", None) and _pt(r.font.size)})
            if sizes:
                tags.append(f"{sizes[-1]}pt")
        except Exception:
            pass
        try:
            if p.alignment == WD_ALIGN_PARAGRAPH.CENTER:
                tags.append("căn giữa")
        except Exception:
            pass
        suffix = f" [{', '.join(tags)}]" if tags else ""
        lines.append(f"[đoạn {i + 1}]{suffix} {p.text.strip()[:300]}")
    if len(paras) > max_paras:
        lines.append(f"... (còn {len(paras) - max_paras} đoạn)")
    lines.append("")

    lines.append(f"## Bảng ({len(doc.tables)} bảng)")
    for ti, table in enumerate(doc.tables):
        lines.append(f"### Bảng {ti + 1} ({len(table.rows)} hàng x "
                     f"{len(table.columns)} cột)")
        for ri, row in enumerate(table.rows[:max_rows]):
            cells = [" ".join(p.text.strip() for p in cell.paragraphs).strip()[:200]
                     for cell in row.cells]
            lines.append(f"hàng {ri + 1}: {' | '.join(cells)}")
        if len(table.rows) > max_rows:
            lines.append(f"... (còn {len(table.rows) - max_rows} hàng)")
    lines.append("")

    for si, sec in enumerate(doc.sections):
        htext, ftext = "", ""
        try:
            htext = _part_text(sec.header)
        except Exception:
            pass
        try:
            ftext = _part_text(sec.footer)
        except Exception:
            pass
        if htext or ftext:
            lines.append(f"## Header/Footer section {si + 1}")
            if htext:
                lines.append(f"header: {htext[:500]}")
            if ftext:
                lines.append(f"footer: {ftext[:500]}")
    lines.append("")

    n_img = sum(1 for _ in _iter_images(doc))
    lines.append(f"## Ảnh ({n_img} ảnh trong toàn file)")
    if logo_rel:
        lines.append(f"- ảnh header đã bóc: {logo_rel}")
    lines.append("")
    lines.append(f"Tổng: {len(paras)} đoạn, {len(doc.tables)} bảng, {n_img} ảnh.")
    return "\n".join(lines)


def _iter_images(doc: Document):
    """Yield every <blip> image element: body + all headers/footers."""
    for el in doc.element.iter():
        if str(getattr(el, "tag", "")).endswith("}blip"):
            yield el
    for sec in doc.sections:
        for part in (sec.header, sec.footer):
            try:
                element = part._element
            except Exception:
                continue
            for el in element.iter():
                if str(getattr(el, "tag", "")).endswith("}blip"):
                    yield el


def extract_visual(docx_path: Path) -> dict:
    """Extract style-level visual params from a .docx file."""
    doc = Document(str(docx_path))
    info: dict = {
        "body_font": "",
        "body_size": None,
        "body_line_spacing": None,
        "body_space_before": None,
        "body_space_after": None,
        "h1_size": None,
        "h2_size": None,
        "h3_size": None,
        "h4_size": None,
        "h5_size": None,
        "h6_size": None,
        "heading_space_before": None,
        "heading_space_after": None,
        "bullet_indent_cm": None,
        "bullet_space_after": None,
        "code_font": "",
        "code_size": None,
        "page_size": "",
        "page_orientation": "",
        "page_margins_cm": {},
        "page_header_distance_cm": None,
        "page_footer_distance_cm": None,
        "has_header": False,
        "header_text": "",
        "has_footer": False,
        "footer_text": "",
        "header_has_image": False,
    }
    normal = _style_font(doc, "Normal")
    if normal.get("name"):
        info["body_font"] = normal["name"]
    if normal.get("size"):
        info["body_size"] = normal["size"]
    try:
        pf = _style_paragraph(doc.styles["Normal"])
        info["body_line_spacing"] = pf.get("line_spacing")
        info["body_space_before"] = pf.get("space_before")
        info["body_space_after"] = pf.get("space_after")
    except KeyError:
        pass
    for key, style_name in (("h1_size", "Heading 1"), ("h2_size", "Heading 2"),
                            ("h3_size", "Heading 3"), ("h4_size", "Heading 4"),
                            ("h5_size", "Heading 5"), ("h6_size", "Heading 6")):
        s = _style_font(doc, style_name)
        if s.get("size"):
            info[key] = s["size"]
        if s.get("name") and not info["body_font"]:
            info["body_font"] = s["name"]
    try:
        hpf = _style_paragraph(doc.styles["Heading 1"])
        info["heading_space_before"] = hpf.get("space_before")
        info["heading_space_after"] = hpf.get("space_after")
    except KeyError:
        pass
    for list_style in ("List Bullet", "List Number"):
        try:
            lpf = _style_paragraph(doc.styles[list_style])
        except KeyError:
            continue
        if lpf.get("indent_cm") and not info["bullet_indent_cm"]:
            info["bullet_indent_cm"] = lpf["indent_cm"]
        if lpf.get("space_after") is not None and info["bullet_space_after"] is None:
            info["bullet_space_after"] = lpf["space_after"]
        if info["bullet_indent_cm"] is not None and info["bullet_space_after"] is not None:
            break
    code = _style_font(doc, "Code Block")
    if code.get("name"):
        info["code_font"] = code["name"]
    if code.get("size"):
        info["code_size"] = code["size"]
    try:
        doc.styles["Code Block"]
        info["has_code_style"] = True
    except KeyError:
        info["has_code_style"] = False

    if doc.sections:
        sec = doc.sections[0]
        try:
            w, h = _cm(sec.page_width), _cm(sec.page_height)
            if w and h:
                if abs(w - 21.0) < 0.15 and abs(h - 29.7) < 0.15:
                    info["page_size"] = "A4"
                elif abs(w - 21.59) < 0.15 and abs(h - 27.94) < 0.15:
                    info["page_size"] = "Letter"
                else:
                    info["page_size"] = f"{w}x{h}"
                info["page_orientation"] = "portrait" if h >= w else "landscape"
            info["page_margins_cm"] = {
                "top": _cm(sec.top_margin),
                "bottom": _cm(sec.bottom_margin),
                "left": _cm(sec.left_margin),
                "right": _cm(sec.right_margin),
            }
            info["page_header_distance_cm"] = _cm(sec.header_distance)
            info["page_footer_distance_cm"] = _cm(sec.footer_distance)
        except Exception:
            pass
    for sec in doc.sections:
        try:
            htext = _part_text(sec.header)
            himg = _has_images_in_element(sec.header._element)
            if htext or himg:
                info["has_header"] = True
                info["header_text"] = info["header_text"] or htext
                info["header_has_image"] = info["header_has_image"] or himg
        except Exception:
            pass
        try:
            ftext = _part_text(sec.footer)
            if ftext:
                info["has_footer"] = True
                info["footer_text"] = info["footer_text"] or ftext
        except Exception:
            pass
    return info


def load_defaults() -> dict:
    if DEFAULTS_YAML.exists():
        with open(DEFAULTS_YAML, encoding="utf-8") as f:
            data = yaml.safe_load(f) or {}
        return data if isinstance(data, dict) else {}
    return {}


def build_template_yaml(tid: str, location: str, source: Path, name: str,
                         description: str, visual: dict,
                         logo_rel: str = "") -> dict:
    base = load_defaults()
    styles = dict(base.get("styles", {}))
    if visual.get("body_font"):
        styles["body_font"] = visual["body_font"]
    for vkey in ("body_size", "body_line_spacing", "body_space_before", "body_space_after",
                 "h1_size", "h2_size", "h3_size", "h4_size", "h5_size", "h6_size",
                 "heading_space_before", "heading_space_after",
                 "bullet_indent_cm", "bullet_space_after"):
        if visual.get(vkey) is not None:
            styles[vkey] = visual[vkey]
    if visual.get("code_font"):
        styles["code_font"] = visual["code_font"]
    if visual.get("code_size"):
        styles["code_size"] = visual["code_size"]

    page = dict(base.get("page", {}))
    if visual.get("page_size") in ("A4", "Letter"):
        page["size"] = visual["page_size"]
    if visual.get("page_orientation") in ("portrait", "landscape"):
        page["orientation"] = visual["page_orientation"]
    margins = visual.get("page_margins_cm") or {}
    for side, pkey in (("top", "margin_top_cm"), ("bottom", "margin_bottom_cm"),
                       ("left", "margin_left_cm"), ("right", "margin_right_cm")):
        if margins.get(side) is not None:
            page[pkey] = margins[side]
    if visual.get("page_header_distance_cm") is not None:
        page["header_distance_cm"] = visual["page_header_distance_cm"]
    if visual.get("page_footer_distance_cm") is not None:
        page["footer_distance_cm"] = visual["page_footer_distance_cm"]

    now = now_iso()
    cfg = dict(base)
    cfg.update(
        {
            "id": tid,
            "name": name or tid,
            "description": description or "",
            "location": location,
            "created_at": now,
            "updated_at": now,
            "source_file": source.name,
            "styles": styles,
            "page": page,
            "_extracted": {
                "has_header": visual.get("has_header", False),
                "header_text": visual.get("header_text", ""),
                "header_has_image": visual.get("header_has_image", False),
                "has_footer": visual.get("has_footer", False),
                "footer_text": visual.get("footer_text", ""),
                "has_code_style": visual.get("has_code_style", False),
                "logo_saved": logo_rel,
                "note": "Khối _extracted chỉ để --validate đối chiếu, không dùng khi render. "
                        "Mục đích mẫu và giá trị params do agent đọc content.txt rồi đề xuất, "
                        "user duyệt mới điền (xem commands/create.md).",
            },
        }
    )
    return cfg


def parse_args(argv=None):
    ap = argparse.ArgumentParser(description="haro-docx template creator")
    ap.add_argument("--source", required=True, help="nguon .docx")
    ap.add_argument("--id", required=True, help="id mau (vd congty-a)")
    ap.add_argument("--location", required=True, choices=["local", "global"])
    ap.add_argument("--project-root", default=".")
    ap.add_argument("--name", default="")
    ap.add_argument("--description", default="")
    return ap.parse_args(argv)


def main(argv=None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="backslashreplace")
        except Exception:
            pass
    args = parse_args(argv)
    try:
        tid = normalize_id(args.id)
    except ValueError as e:
        print(str(e), file=sys.stderr)
        return 2
    source = Path(args.source)
    if not source.exists():
        print(f"LỖI: không tìm thấy file '{source}'.", file=sys.stderr)
        return 2
    if source.suffix.lower() != ".docx":
        print("LỖI: file tạo mẫu phải là .docx.", file=sys.stderr)
        return 2
    root = local_root(Path(args.project_root)) if args.location == "local" else global_root()
    if exists_in(root, tid):
        print(
            f"LỖI: id '{tid}' đã tồn tại ở {args.location} ({root / tid}). "
            f"Dùng /haro-docx --view:{tid} để xem hoặc --update:{tid} để sửa.",
            file=sys.stderr,
        )
        return 3
    paths = template_paths(root, tid)
    try:
        doc = Document(str(source))
        visual = extract_visual(source)
        paths["dir"].mkdir(parents=True, exist_ok=True)
        logo_rel = save_logo_blob(_header_image_blobs(doc), paths["dir"])
    except Exception as e:
        print(f"LỖI: không đọc được file .docx ({e}).", file=sys.stderr)
        return 2
    copy_source_into(source, paths["docx"])
    cfg = build_template_yaml(tid, args.location, source, args.name,
                              args.description, visual, logo_rel)
    with open(paths["yaml"], "w", encoding="utf-8") as f:
        yaml.safe_dump(cfg, f, allow_unicode=True, sort_keys=False)
    content_txt = paths["dir"] / "content.txt"
    with open(content_txt, "w", encoding="utf-8") as f:
        f.write(dump_content(doc, source.name, logo_rel))
    st = cfg["styles"]
    print(f"Đã tạo mẫu '{tid}' ở {args.location}.")
    print(f"File mẫu: {paths['docx']}")
    print(f"File YAML: {paths['yaml']}")
    print("Style trích xuất (sửa trong YAML nếu cần):")
    print(f"  body: font={st.get('body_font')} size={st.get('body_size')} "
          f"line={st.get('body_line_spacing')} space={st.get('body_space_before')}/{st.get('body_space_after')}")
    print(f"  h1-h6: {st.get('h1_size')}/{st.get('h2_size')}/{st.get('h3_size')}/"
          f"{st.get('h4_size')}/{st.get('h5_size')}/{st.get('h6_size')}")
    print(f"  bullet: indent={st.get('bullet_indent_cm')}cm space_after={st.get('bullet_space_after')}pt")
    pg = cfg.get("page", {})
    print(f"  trang: {pg.get('size')} {pg.get('orientation')} "
          f"margins={pg.get('margin_top_cm')}/{pg.get('margin_bottom_cm')}/"
          f"{pg.get('margin_left_cm')}/{pg.get('margin_right_cm')}cm")
    print(f"  header: {'có' if visual.get('has_header') else 'không có'}; "
          f"footer: {'có' if visual.get('has_footer') else 'không có'}"
          + (f"; logo đã bóc: {logo_rel}" if logo_rel else ""))
    print(f"Nội dung mẫu (agent đọc để hiểu mục đích): {content_txt}")
    print("Tiếp theo: đọc content.txt, nêu mục đích mẫu, đề xuất thông số để user duyệt.")
    print(f"Kiểm tra lại: /haro-docx --view:{tid}  |  Xuất thử: /haro-docx --export:{tid} <file.md>")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
