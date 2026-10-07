#!/usr/bin/env python3
"""Apply YAML param changes to a haro-docx-writer template (yaml + .docx styles).

Typical flow for `/haro-docx-writer --update:<id> <nội dung>`: the agent edits
template.yaml (or passes --set key=value here), then this script pushes
style-level keys into template.docx via python-docx and bumps updated_at.

Usage:
    python update_template.py --id <template-id> [--project-root .]
        [--set styles.h1_size=15 --set company_name="CTY X"]
        [--no-apply-visual]

--set accepts dot-paths: name, description, company_name, solution_name,
document_name, document_title, version, date, status, header.logo_path,
styles.body_font, styles.body_size, styles.h1_size ... styles.h6_size,
styles.body_line_spacing, styles.body_space_before/after,
styles.heading_space_before/after, styles.bullet_indent_cm,
styles.bullet_space_after, styles.code_font, styles.code_size,
styles.heading_color, page.size/orientation/margins/distances,
cover.*_size, header_footer.*_size, toc_levels,
placeholders.<tên>.value | placeholders.<tên>.description (free text).
authors/reviewers/approvers accept comma-separated lists.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import List, Optional

import yaml
from docx import Document
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from template_store import (  # noqa: E402
    fmt_unit,
    normalize_id,
    now_iso,
    parse_cm,
    parse_pt,
    resolve,
)

STYLE_KEYS = {
    "styles.body_font", "styles.body_size",
    "styles.h1_size", "styles.h2_size", "styles.h3_size",
    "styles.h4_size", "styles.h5_size", "styles.h6_size",
    "styles.body_line_spacing", "styles.body_space_before", "styles.body_space_after",
    "styles.heading_space_before", "styles.heading_space_after",
    "styles.bullet_indent_cm", "styles.bullet_space_after",
    "styles.code_font", "styles.code_size", "styles.heading_color",
}
PAGE_KEYS = {
    "page.size", "page.orientation",
    "page.margin_top_cm", "page.margin_bottom_cm",
    "page.margin_left_cm", "page.margin_right_cm",
    "page.header_distance_cm", "page.footer_distance_cm",
}
COVER_KEYS = {
    "cover.title_size", "cover.solution_size", "cover.company_size", "cover.note_size",
}
HF_KEYS = {
    "header_footer.doc_name_size", "header_footer.company_size",
    "header_footer.page_size", "header_footer.solution_size",
}
TEXT_KEYS = {
    "name", "description", "company_name", "solution_name", "document_name",
    "document_title", "version", "date", "status", "header.logo_path",
    "toc_levels",
}
LIST_KEYS = {"authors", "reviewers", "approvers"}
ALLOWED = STYLE_KEYS | PAGE_KEYS | COVER_KEYS | HF_KEYS | TEXT_KEYS | LIST_KEYS
# Keys measured in points (font/spacing sizes) vs centimeters (page geometry).
PT_KEYS = {
    "styles.body_size", "styles.h1_size", "styles.h2_size", "styles.h3_size",
    "styles.h4_size", "styles.h5_size", "styles.h6_size",
    "styles.body_space_before", "styles.body_space_after",
    "styles.heading_space_before", "styles.heading_space_after",
    "styles.bullet_space_after",
    "styles.code_size",
    "cover.title_size", "cover.solution_size", "cover.company_size", "cover.note_size",
    "header_footer.doc_name_size", "header_footer.company_size",
    "header_footer.page_size", "header_footer.solution_size",
}
CM_KEYS = {
    "styles.bullet_indent_cm",
    "page.margin_top_cm", "page.margin_bottom_cm",
    "page.margin_left_cm", "page.margin_right_cm",
    "page.header_distance_cm", "page.footer_distance_cm",
}
NUMERIC_KEYS = PT_KEYS | CM_KEYS | {"styles.body_line_spacing"}


def _parse_value(key: str, raw: str):
    if key in LIST_KEYS:
        return [p.strip() for p in raw.split(",") if p.strip()]
    if _is_placeholder_key(key):
        return raw  # placeholder value/description: free text verbatim
    if key in PT_KEYS:
        # Stored with unit suffix for readability (vd "12pt").
        try:
            val = parse_pt(raw, key)
        except ValueError as e:
            raise ValueError(str(e))
        return fmt_unit(val, "pt")
    if key in CM_KEYS:
        try:
            val = parse_cm(raw, key)
        except ValueError as e:
            raise ValueError(str(e))
        return fmt_unit(val, "cm")
    if key == "styles.body_line_spacing":
        try:
            num = float(str(raw).strip())
        except ValueError:
            raise ValueError(f"LỖI: '{key}' phải là hệ số (vd 1.15), nhận được '{raw}'.")
        if num <= 0:
            raise ValueError(f"LỖI: '{key}' phải > 0 (nhận được '{raw}').")
        return int(num) if num.is_integer() else num
    if key == "page.size":
        v = raw.strip().upper()
        if v not in ("A4", "LETTER"):
            raise ValueError(f"LỖI: 'page.size' chỉ nhận A4 | Letter (nhận được '{raw}').")
        return v
    if key == "page.orientation":
        v = raw.strip().lower()
        if v not in ("portrait", "landscape"):
            raise ValueError(f"LỖI: 'page.orientation' chỉ nhận portrait | landscape (nhận được '{raw}').")
        return v
    if key == "toc_levels":
        import re as _re
        v = raw.strip()
        if not _re.fullmatch(r"1-[1-9]", v):
            raise ValueError(f"LỖI: 'toc_levels' dạng '1-N' (vd 1-6), nhận được '{raw}'.")
        return v
    if key == "styles.heading_color":
        v = raw.strip().lstrip("#")
        import re as _re
        if not _re.fullmatch(r"[0-9a-fA-F]{6}", v):
            raise ValueError(f"LỖI: 'styles.heading_color' là hex 6 ký tự (vd 000000), nhận được '{raw}'.")
        return v.upper()
    return raw


def _is_placeholder_key(key: str) -> bool:
    """placeholders.<name>.value | placeholders.<name>.description only."""
    parts = key.split(".")
    return (
        len(parts) >= 3
        and parts[0] == "placeholders"
        and parts[-1] in ("value", "description")
        and all(p for p in parts[1:-1])
    )


def _get(cfg: dict, key: str):
    parts = key.split(".")
    cur: object = cfg
    for p in parts:
        if not isinstance(cur, dict) or p not in cur:
            return None
        cur = cur[p]
    return cur


def _set(cfg: dict, key: str, value) -> None:
    parts = key.split(".")
    cur = cfg
    for p in parts[:-1]:
        if p not in cur or not isinstance(cur[p], dict):
            cur[p] = {}
        cur = cur[p]
    cur[parts[-1]] = value


def apply_visual(docx_path: Path, styles: dict, page: Optional[dict] = None) -> List[str]:
    """Push style-level params into the .docx file. Returns changed labels."""
    doc = Document(str(docx_path))
    changed: List[str] = []
    body_font = styles.get("body_font", "")
    code_font = styles.get("code_font", "")

    def touch(style_name: str, size_key: str, font: str = ""):
        try:
            st = doc.styles[style_name]
        except KeyError:
            return
        size = styles.get(size_key)
        if size is not None and size != "":
            st.font.size = Pt(parse_pt(size, f"styles.{size_key}"))
            changed.append(f"{style_name}.size={size}")
        if font:
            try:
                st.font.name = font
                changed.append(f"{style_name}.font={font}")
            except Exception:
                pass

    if body_font or styles.get("body_size"):
        touch("Normal", "body_size", body_font)
    try:
        npf = doc.styles["Normal"].paragraph_format
        if styles.get("body_line_spacing") is not None:
            npf.line_spacing = float(str(styles["body_line_spacing"]).strip())
            changed.append(f"Normal.line_spacing={styles['body_line_spacing']}")
        if styles.get("body_space_before") is not None:
            npf.space_before = Pt(parse_pt(styles["body_space_before"], "styles.body_space_before"))
            changed.append(f"Normal.space_before={styles['body_space_before']}")
        if styles.get("body_space_after") is not None:
            npf.space_after = Pt(parse_pt(styles["body_space_after"], "styles.body_space_after"))
            changed.append(f"Normal.space_after={styles['body_space_after']}")
    except KeyError:
        pass
    for name, skey in (("Heading 1", "h1_size"), ("Heading 2", "h2_size"),
                       ("Heading 3", "h3_size"), ("Heading 4", "h4_size"),
                       ("Heading 5", "h5_size"), ("Heading 6", "h6_size")):
        touch(name, skey, body_font)
        try:
            hpf = doc.styles[name].paragraph_format
            if styles.get("heading_space_before") is not None:
                hpf.space_before = Pt(parse_pt(styles["heading_space_before"],
                                               "styles.heading_space_before"))
            if styles.get("heading_space_after") is not None:
                hpf.space_after = Pt(parse_pt(styles["heading_space_after"],
                                              "styles.heading_space_after"))
        except KeyError:
            pass
    if styles.get("heading_space_before") is not None or styles.get("heading_space_after") is not None:
        changed.append(f"heading_space={styles.get('heading_space_before')}/{styles.get('heading_space_after')}")
    for list_name in ("List Bullet", "List Number"):
        try:
            lst = doc.styles[list_name]
        except KeyError:
            continue
        if styles.get("bullet_indent_cm") is not None:
            lst.paragraph_format.left_indent = Cm(parse_cm(styles["bullet_indent_cm"],
                                                          "styles.bullet_indent_cm"))
            changed.append(f"{list_name}.indent={styles['bullet_indent_cm']}")
        if styles.get("bullet_space_after") is not None:
            lst.paragraph_format.space_after = Pt(parse_pt(styles["bullet_space_after"],
                                                           "styles.bullet_space_after"))
            changed.append(f"{list_name}.space_after={styles['bullet_space_after']}")
    if "Code Block" in [s.name for s in doc.styles]:
        touch("Code Block", "code_size", code_font)

    if page:
        size = str(page.get("size", "A4")).upper()
        w_cm, h_cm = (21.59, 27.94) if size == "LETTER" else (21.0, 29.7)
        if str(page.get("orientation", "portrait")).lower() == "landscape":
            w_cm, h_cm = h_cm, w_cm
        for s in doc.sections:
            s.page_width = Cm(w_cm)
            s.page_height = Cm(h_cm)
            for attr, pkey in (("top_margin", "margin_top_cm"), ("bottom_margin", "margin_bottom_cm"),
                               ("left_margin", "margin_left_cm"), ("right_margin", "margin_right_cm"),
                               ("header_distance", "header_distance_cm"),
                               ("footer_distance", "footer_distance_cm")):
                if page.get(pkey) is not None:
                    setattr(s, attr, Cm(parse_cm(page[pkey], f"page.{pkey}")))
        changed.append(f"page={size} {page.get('orientation', 'portrait')} "
                       f"margins={page.get('margin_top_cm')}/{page.get('margin_bottom_cm')}/"
                       f"{page.get('margin_left_cm')}/{page.get('margin_right_cm')}")
    doc.save(str(docx_path))
    return changed


def parse_args(argv=None):
    ap = argparse.ArgumentParser(description="haro-docx-writer template updater")
    ap.add_argument("--id", required=True)
    ap.add_argument("--project-root", default=".")
    ap.add_argument("--set", action="append", default=[], help="key=value (dot-path)")
    ap.add_argument("--no-apply-visual", action="store_true",
                    help="chi sua yaml, khong day style vao docx")
    return ap.parse_args(argv)


def main(argv=None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="backslashreplace")
        except Exception:
            pass
    args = parse_args(argv)
    try:
        tid = normalize_id(args.id)
    except ValueError as e:
        print(str(e), file=sys.stderr)
        return 2
    hit = resolve(tid, Path(args.project_root))
    if not hit or not hit["yaml"].exists():
        print(f"LỖI: không tìm thấy mẫu '{tid}'.", file=sys.stderr)
        return 2
    with open(hit["yaml"], encoding="utf-8") as f:
        cfg = yaml.safe_load(f) or {}

    diffs: List[str] = []
    for item in args.set:
        if "=" not in item:
            print(f"LỖI: --set phải dạng key=value (nhận được '{item}').", file=sys.stderr)
            return 2
        key, raw = item.split("=", 1)
        key, raw = key.strip(), raw.strip()
        if key not in ALLOWED and not _is_placeholder_key(key):
            print(f"LỖI: key '{key}' không hỗ trợ. Key hợp lệ: {', '.join(sorted(ALLOWED))} "
                  f"+ placeholders.<tên>.value|.description.",
                  file=sys.stderr)
            return 2
        try:
            value = _parse_value(key, raw)
        except ValueError as e:
            print(str(e), file=sys.stderr)
            return 2
        old = _get(cfg, key)
        _set(cfg, key, value)
        diffs.append(f"{key}: '{old}' -> '{value}'")

    cfg["updated_at"] = now_iso()
    with open(hit["yaml"], "w", encoding="utf-8") as f:
        yaml.safe_dump(cfg, f, allow_unicode=True, sort_keys=False)

    pushed: List[str] = []
    if not args.no_apply_visual and hit["docx"].exists():
        try:
            pushed = apply_visual(hit["docx"], cfg.get("styles", {}) or {},
                                  cfg.get("page", {}) or {})
        except Exception as e:
            print(f"CẢNH BÁO: đã lưu YAML nhưng không đẩy được style vào .docx ({e}).", file=sys.stderr)
    if diffs:
        print(f"Đã cập nhật mẫu '{tid}':")
        for d in diffs:
            print(f"  - {d}")
    else:
        print(f"Đã chạm updated_at cho mẫu '{tid}' (không có --set).")
    if pushed:
        print("Đã đẩy style vào file mẫu:")
        for p in pushed:
            print(f"  - {p}")
    print(f"YAML: {hit['yaml']}")
    print(f"DOCX: {hit['docx']}")
    print(f"Kiểm tra: /haro-docx-writer --validate:{tid}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
